import os
import importlib
import pandas as pd
from io import StringIO, BytesIO
from mistral_ai import traer_texto_png
from connect_gemini import estructurar_con_prompt_imgia, limpiar_csv_de_respuesta
from google.oauth2 import service_account
from googleapiclient.discovery import build
from openpyxl import load_workbook
from PIL import Image
import requests
import csv
from google.api_core.client_options import ClientOptions
from google.cloud import documentai_v1 as documentai

# --- Configuración Google Drive ---
SERVICE_ACCOUNT_FILE = 'credentials/credentials.json'
SCOPES = ['https://www.googleapis.com/auth/drive.readonly']
creds = service_account.Credentials.from_service_account_file(SERVICE_ACCOUNT_FILE, scopes=SCOPES)
service = build('drive', 'v3', credentials=creds)

# --- Configuración Document AI ---
os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = "google_ocr.json"
PROJECT_ID = "600456283474"
LOCATION = "us"
PROCESSOR_ID = "b668a0883f50e7fa"
ENDPOINT = "us-documentai.googleapis.com"
PROCESSOR_NAME = f"projects/{PROJECT_ID}/locations/{LOCATION}/processors/{PROCESSOR_ID}"
client_options = ClientOptions(api_endpoint=ENDPOINT)
docai_client = documentai.DocumentProcessorServiceClient(client_options=client_options)

# --- Configuración general ---
proveedores_documentai = ["quilmes", "pepita"]
columnas_img = ["Código Gem", "Producto Gem", "Cantidad Gem", "Precio Gem", "Total Gem"]

# --- Funciones auxiliares ---
def listar_imagenes(folder_id):
    query = f"'{folder_id}' in parents and trashed = false"
    resultados = service.files().list(q=query, fields="files(id, name, mimeType)").execute()
    return [f for f in resultados.get('files', []) if f['mimeType'].startswith('image/')]

def escribir_en_bd_local(df_resultado):
    ruta_archivo = "C:/Users/gesti/Desktop/extractor_de_facturas_v6/BD extracciones.xlsx"
    try:
        libro = load_workbook(ruta_archivo)
        hoja = libro.active
        for _, fila in df_resultado.iterrows():
            hoja.append(fila.tolist())
        libro.save(ruta_archivo)
        print(f"\n✅ Datos agregados correctamente en '{ruta_archivo}'")
    except Exception as e:
        print(f"\n❌ Error al escribir en '{ruta_archivo}': {e}")

def obtener_imagen_desde_drive(file_id):
    url = f"https://drive.google.com/uc?id={file_id}&export=download"
    response = requests.get(url)
    response.raise_for_status()
    return Image.open(BytesIO(response.content))

def procesar_con_document_ai(file_id):
    try:
        url = f"https://drive.google.com/uc?id={file_id}&export=download"
        response = requests.get(url)
        response.raise_for_status()
        image_bytes = response.content

        raw_document = documentai.RawDocument(content=image_bytes, mime_type="image/jpeg")
        request = documentai.ProcessRequest(name=PROCESSOR_NAME, raw_document=raw_document)
        result = docai_client.process_document(request=request)
        return result.document

    except Exception as e:
        print(f"❌ Error al procesar imagen con Document AI: {e}")
        return None

def alinear_y_combinar(df_ocr, df_img_raw):
    # Si IMGIA no devolvió nada, devolvemos solo OCRIA (aunque esté vacío)
    if df_img_raw.empty:
        return df_ocr

    # Si OCRIA está vacío pero IMGIA tiene datos, rellenamos OCRIA con filas vacías para que coincidan
    if df_ocr.empty:
        df_ocr = pd.DataFrame([[""] * len(df_img_raw)] * len(df_img_raw), columns=[f"OCR_Col_{i}" for i in range(len(df_img_raw))])
        return pd.concat([df_ocr.reset_index(drop=True), df_img_raw.reset_index(drop=True)], axis=1)

    # Si OCRIA tiene solo una fila y IMGIA tiene varias, replicamos la fila de OCRIA
    if len(df_ocr) == 1 and len(df_img_raw) > 1:
        df_ocr = pd.DataFrame([df_ocr.iloc[0].tolist()] * len(df_img_raw), columns=df_ocr.columns)

    # Si tienen distinta cantidad de filas, reindexamos IMGIA para que coincida con OCRIA
    elif len(df_ocr) != len(df_img_raw):
        df_img_raw = df_img_raw.reindex(range(len(df_ocr))).fillna("")

    return pd.concat([df_ocr.reset_index(drop=True), df_img_raw.reset_index(drop=True)], axis=1)


def main():
    folder_id = '1YCEgOfDfyD9levr4_ouoXpxy0l0n5QXH'
    imagenes = listar_imagenes(folder_id)
    resultados_df = []

    for img in imagenes:
        nombre_archivo = img['name'].lower()
        file_id = img['id']

        for archivo in os.listdir("proveedores"):
            if archivo.endswith(".py") and archivo != "__init__.py":
                proveedor = archivo[:-3].lower()
                if proveedor in nombre_archivo:
                    print(f"\n📄 Procesando imagen: {img['name']} con proveedor: {proveedor}")
                    modulo = importlib.import_module(f'proveedores.{proveedor}')

                    if proveedor in proveedores_documentai:
                        print("🔍 OCR con Document AI")
                        document = procesar_con_document_ai(file_id)
                        if not document:
                            continue
                        try:
                            procesar = getattr(modulo, "procesar")
                            df_ocr = procesar(document, img['name'])
                        except Exception as e:
                            print(f"❌ Error procesando proveedor '{proveedor}': {e}")
                            continue
                        df_final = df_ocr

                    else:
                        texto = traer_texto_png(file_id)
                        if not texto:
                            print("❌ No se pudo extraer texto con Mistral OCR.")
                            continue

                        print("📤 Texto OCR Mistral:")
                        print(texto)

                        df_ocr = pd.DataFrame()
                        try:
                            df_ocr = modulo.procesar(texto)
                            if df_ocr is None or df_ocr.empty:
                                print(f"⚠️ No se pudo procesar OCRIA para proveedor '{proveedor}', se continuará con IMGIA.")
                                df_ocr = pd.DataFrame()
                            else:
                                print("\n✅ Resultado estructurado OCRIA:")
                                print(df_ocr)
                        except Exception as e:
                            print(f"⚠️ Error procesando OCRIA para proveedor '{proveedor}': {e}. Se continuará con IMGIA.")
                            df_ocr = pd.DataFrame()

                        # Gemini IMGIA
                        df_img_raw = pd.DataFrame(columns=columnas_img)
                        download_url = f"https://drive.google.com/uc?id={file_id}&export=download"

                        try:
                            prompt_img = modulo.prompt_imgia(download_url)
                            imagen_pil = obtener_imagen_desde_drive(file_id)

                            resultado_img = estructurar_con_prompt_imgia(prompt_img, imagen_pil)
                            resultado_img = limpiar_csv_de_respuesta(resultado_img)
                            print(f"\n🧾 Texto crudo devuelto por IMGIA:\n{resultado_img}")

                            if resultado_img.strip() == "":
                                print("⚠️ IMGIA devolvió texto vacío.")
                                df_img_raw = pd.DataFrame(columns=columnas_img)
                            else:
                                try:
                                    df_img_raw = pd.read_csv(
                                        StringIO(resultado_img),
                                        header=None,
                                        skip_blank_lines=True,
                                        on_bad_lines='warn',
                                        quotechar='"',
                                        sep=","
                                    )
                                    print(f"📊 IMGIA devolvió {df_img_raw.shape[0]} filas y {df_img_raw.shape[1]} columnas")

                                    if df_img_raw.shape[1] == len(columnas_img):
                                        df_img_raw.columns = columnas_img
                                        for col, func_name in [
                                            ("Cantidad Gem", "limpiar_cantidad"),
                                            ("Precio Gem", "limpiar_numero"),
                                            ("Total Gem", "limpiar_numero")
                                        ]:
                                            try:
                                                func = getattr(modulo, func_name)
                                                df_img_raw[col] = df_img_raw[col].apply(func)
                                            except AttributeError:
                                                print(f"⚠️ {func_name} no está definido en el módulo '{proveedor}'. Se deja sin aplicar.")
                                            except Exception as e:
                                                print(f"⚠️ Error aplicando {func_name}: {e}")
                                    else:
                                        print(f"⚠️ CSV IMGIA mal formado. Columnas: {df_img_raw.shape[1]}")
                                        df_img_raw = pd.DataFrame(columns=columnas_img)
                                except Exception as e:
                                    print(f"❌ Error leyendo CSV de IMGIA: {e}")
                                    df_img_raw = pd.DataFrame(columns=columnas_img)

                        except AttributeError:
                            print(f"⚠️ El módulo '{proveedor}' no tiene la función 'prompt_imgia'. Se salta IMGIA.")
                        except Exception as e:
                            print(f"❌ Error procesando IMGIA para '{proveedor}': {e}")

                        df_final = alinear_y_combinar(df_ocr, df_img_raw)

                    # Validar columnas duplicadas
                    if pd.Series(df_final.columns).duplicated().any():
                        df_final.columns = [
                            f"{col}_{i}" if pd.Series(df_final.columns).duplicated()[j] else col
                            for j, (i, col) in enumerate(zip(range(len(df_final.columns)), df_final.columns))
                        ]

                    # --- Asegurar columnas Gem primero y orden final ---
                    columnas_gem = ["Código Gem", "Producto Gem", "Cantidad Gem", "Precio Gem", "Total Gem"]
                    columnas_ocr = ["Codigo", "Producto", "Cantidad", "Precio", "Total", "Local", "Proveedor"]

                    for col in columnas_gem:
                        if col not in df_final.columns:
                            df_final[col] = ""

                    for col in columnas_ocr:
                        if col not in df_final.columns:
                            df_final[col] = ""

                    df_final["Nombre Archivo"] = img['name']
                    columnas_ordenadas = columnas_gem + columnas_ocr + ["Nombre Archivo"]
                    df_final = df_final[[col for col in columnas_ordenadas if col in df_final.columns]]

                    resultados_df.append(df_final)
                    print(f"✅ Resultado final para {proveedor}:")
                    print(df_final)
                    break
        else:
            print(f"❌ No se encontró proveedor coincidente para: {nombre_archivo}")

    if resultados_df:
        df_total = pd.concat(resultados_df, ignore_index=True)
        escribir_en_bd_local(df_total)
    else:
        print("⚠️ No se generaron resultados para guardar.")



if __name__ == "__main__":
    main()
